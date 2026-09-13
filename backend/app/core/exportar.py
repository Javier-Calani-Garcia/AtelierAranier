import csv
import io
from typing import Any

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BRAND_DARK = "203C40"


def exportar_csv(nombre_archivo: str, headers: list[str], filas: list[list[Any]]) -> Response:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(filas)
    return Response(
        content=buffer.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}.csv"'},
    )


def exportar_excel(nombre_archivo: str, titulo: str, headers: list[str], filas: list[list[Any]]) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
    celda_titulo = ws.cell(row=1, column=1, value=f"Atelier Aranier -- {titulo}")
    celda_titulo.font = Font(bold=True, size=14, color=BRAND_DARK)
    ws.row_dimensions[1].height = 24

    fila_headers = 3
    for col, encabezado in enumerate(headers, start=1):
        celda = ws.cell(row=fila_headers, column=col, value=encabezado)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(start_color=BRAND_DARK, end_color=BRAND_DARK, fill_type="solid")
        celda.alignment = Alignment(horizontal="left")

    for row_idx, fila in enumerate(filas, start=fila_headers + 1):
        for col_idx, valor in enumerate(fila, start=1):
            ws.cell(row=row_idx, column=col_idx, value=valor)

    for col in range(1, len(headers) + 1):
        letra = get_column_letter(col)
        largo = max([len(str(headers[col - 1]))] + [len(str(f[col - 1])) for f in filas] + [10])
        ws.column_dimensions[letra].width = min(largo + 2, 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}.xlsx"'},
    )
